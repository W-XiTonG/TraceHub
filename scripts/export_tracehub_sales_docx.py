# -*- coding: utf-8 -*-
"""Convert docs/TraceHub-销售介绍.md to Word .docx (requires: pip install python-docx)."""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:
    print("Missing dependency: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def set_doc_font(doc: Document) -> None:
    """Prefer East Asian font for Chinese on Windows Word."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Microsoft YaHei"
    font.size = Pt(11)
    r = style._element.rPr.rFonts
    if r is not None:
        r.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_mixed_bold(paragraph, text: str) -> None:
    """Split **bold** segments into runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**") and len(p) >= 4:
            run = paragraph.add_run(p[2:-2])
            run.bold = True
        else:
            paragraph.add_run(p)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def parse_table_row(line: str) -> list[str]:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_separator_row(cells: list[str]) -> bool:
    if not cells:
        return False
    return all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells if c)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    md_path = root / "docs" / "TraceHub-销售介绍.md"
    out_path = root / "docs" / "TraceHub-销售介绍.docx"

    if not md_path.is_file():
        print(f"Not found: {md_path}", file=sys.stderr)
        sys.exit(1)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_doc_font(doc)

    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        if raw.strip() == "":
            i += 1
            continue

        if raw.strip() == "---":
            i += 1
            continue

        if is_table_row(raw):
            table_lines: list[str] = []
            while i < len(lines) and is_table_row(lines[i].rstrip()):
                table_lines.append(lines[i].rstrip())
                i += 1
            rows: list[list[str]] = []
            for tl in table_lines:
                cells = parse_table_row(tl)
                if is_separator_row(cells):
                    continue
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncol)
                tbl.style = "Table Grid"
                for ri, row_cells in enumerate(rows):
                    for ci in range(ncol):
                        cell_text = row_cells[ci] if ci < len(row_cells) else ""
                        p = tbl.rows[ri].cells[ci].paragraphs[0]
                        add_mixed_bold(p, cell_text)
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", raw)
        if m:
            hashes = len(m.group(1))
            title = m.group(2).strip()
            # Markdown # / ## / ### → Word Title / Heading 1 / Heading 2
            word_level = {1: 0, 2: 1, 3: 2}.get(hashes, 2)
            p = doc.add_heading(title, level=word_level)
            if word_level == 0:
                p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            i += 1
            continue

        if raw.lstrip().startswith(("- ", "* ")):
            text = raw.lstrip()[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            add_mixed_bold(p, text)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", raw.lstrip()):
            text = re.sub(r"^\d+\.\s+", "", raw.lstrip())
            p = doc.add_paragraph(style="List Number")
            add_mixed_bold(p, text)
            i += 1
            continue

        if raw.lstrip().startswith("*") and raw.rstrip().endswith("*") and not raw.lstrip().startswith("**"):
            # Italic single line like *doc version*
            p = doc.add_paragraph()
            run = p.add_run(raw.strip().strip("*"))
            run.italic = True
            i += 1
            continue

        p = doc.add_paragraph()
        add_mixed_bold(p, raw.strip())
        i += 1

    doc.save(out_path)
    print(f"Wrote: {out_path}")


if __name__ == "__main__":
    main()
